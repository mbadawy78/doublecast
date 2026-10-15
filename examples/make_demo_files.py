"""Create fictional demo files (an Excel P&L and a Word memo) for trying Doublecast.

python examples/make_demo_files.py
doublecast -c examples/config.example.yaml scan examples/demo/customer-review.docx
"""

from pathlib import Path

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font

OUT = Path(__file__).parent / "demo"
OUT.mkdir(exist_ok=True)

wb = Workbook()
ws = wb.active
ws.title = "Q3 customers"
ws.append(["Customer", "Net sales (€)", "Volume (kg)", "Gross margin %", "Account manager"])
rows = [
    ("Nordia Hellas", 1_234_567.89, 410_000, 0.21, "Eleni Papadopoulou"),
    ("Kallergis Markets", 845_300.00, 302_500, 0.24, "Georgios Andreou"),
    ("Brennwald Handel GmbH", 512_940.50, 150_200, 0.19, "Eleni Papadopoulou"),
    ("Harborline Foods", 298_110.00, 88_300, 0.27, "Georgios Andreou"),
]
for r in rows:
    ws.append(list(r))
ws.append(["Total", "=SUM(B2:B5)", "=SUM(C2:C5)", None, None])
for c in ws[1]:
    c.font = Font(bold=True)
for row in ws.iter_rows(min_row=2, max_row=6):
    row[1].number_format = "#,##0.00 €"
    row[3].number_format = "0.0%"
wb.save(OUT / "q3-customers.xlsx")

doc = Document()
doc.add_heading("Q3 customer review: Nordia Hellas", level=1)
p = doc.add_paragraph("Nordia Hellas asked for an extra 8% off invoice no. INV-2026/0418. ")
p.add_run("Eleni Papadopoulou").bold = True
p.add_run(" thinks the account (€1.234.567,89 in Q3) cannot absorb it. Volumes ship from Larissa Plant.")
doc.add_paragraph(
    "Η Καλλέργης ζητά νέα τιμή για το Golden Twist 80g. Επικοινωνία: g.andreou@acmefoods.gr, +30 694 123 4567. "
    "Πληρωμή στο IBAN GR16 0110 1250 0000 0001 2300 695, ΑΦΜ 094014201."
)
doc.add_paragraph("Project Falcon launch stays on schedule.")
doc.save(OUT / "customer-review.docx")

print(f"Wrote {OUT / 'q3-customers.xlsx'} and {OUT / 'customer-review.docx'}")
