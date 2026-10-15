"""A folder of Excel and CSV files (every sheet is a table)."""

from __future__ import annotations

import csv
import io
from collections.abc import Iterator
from pathlib import Path

from .base import Source, Table


def _header_row(rows: list[list]) -> int:
    for i, row in enumerate(rows[:6]):
        if sum(1 for c in row if isinstance(c, str) and c.strip()) >= 2:
            return i
    return 0


class FolderSource(Source):
    kind = "folder"

    def __init__(self, name: str, path: str | Path, recursive: bool = True, max_files: int = 500):
        super().__init__(name)
        self.root = Path(path).expanduser()
        self.recursive = recursive
        self.max_files = max_files
        self._cache: dict[str, list[dict]] = {}

    def _files(self) -> list[Path]:
        if self.root.is_file():
            return [self.root]
        pattern = "**/*" if self.recursive else "*"
        files = [
            p
            for p in self.root.glob(pattern)
            if p.suffix.lower() in (".xlsx", ".xlsm", ".csv", ".tsv")
            and not p.name.startswith(("~$", "."))
            and ".doublecast." not in p.name
            and ".restored." not in p.name
        ]
        return sorted(files)[: self.max_files]

    def tables(self) -> Iterator[Table]:
        for f in self._files():
            if f.suffix.lower() in (".csv", ".tsv"):
                rows = self._csv(f)
                self._cache[str(f)] = rows
                yield Table(str(f), list(rows[0].keys()) if rows else [])
            else:
                from openpyxl import load_workbook

                wb = load_workbook(f, read_only=True, data_only=True)
                for ws in wb.worksheets:
                    raw = [list(r) for r in ws.iter_rows(values_only=True)]
                    if not raw:
                        continue
                    h = _header_row(raw)
                    header = [str(c).strip() if c is not None else f"col{i + 1}" for i, c in enumerate(raw[h])]
                    rows = [dict(zip(header, r, strict=False)) for r in raw[h + 1 :] if any(v is not None for v in r)]
                    key = f"{f}::{ws.title}"
                    self._cache[key] = rows
                    yield Table(key, header)
                wb.close()

    def _csv(self, f: Path) -> list[dict]:
        raw = f.read_text(encoding="utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(raw[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel_tab if f.suffix.lower() == ".tsv" else csv.excel
        return list(csv.DictReader(io.StringIO(raw), dialect=dialect))

    def sample(self, table: Table, n: int) -> list[dict]:
        return self._cache.get(table.name, [])[:n]

    def values(self, table: Table, column: str, limit: int):
        for row in self._cache.get(table.name, [])[:limit]:
            yield row.get(column)
