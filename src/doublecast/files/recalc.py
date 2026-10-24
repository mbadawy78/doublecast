"""Spreadsheet models: recalculate the scrambled copy, and check it still adds up.

openpyxl writes formulas without their results, so an AI reading a scrambled
model with a library sees empty P&L lines. LibreOffice (if installed) fills the
results in.

The check recalculates the original too and compares every formula result:
each must equal the original (a ratio, a count, a %) or the original times the
money factor (an amount). Anything else means the scrambled model no longer
tells the truth, and the file is flagged.
"""

from __future__ import annotations

import os
import shutil
import subprocess
import tempfile
from pathlib import Path

_MAC = "/Applications/LibreOffice.app/Contents/MacOS/soffice"


def find_soffice() -> str | None:
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    return _MAC if os.path.exists(_MAC) else None


def has_formulas(path: Path) -> bool:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True)
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                if any(isinstance(v, str) and v.startswith("=") for v in row):
                    return True
    finally:
        wb.close()
    return False


def recalc(path: Path, timeout: int = 180) -> bool:
    """Recalculate an .xlsx in place with LibreOffice. False if not possible."""
    soffice = find_soffice()
    if not soffice or Path(path).suffix.lower() != ".xlsx":
        return False
    with tempfile.TemporaryDirectory(prefix="dc-lo-") as tmp:
        profile = Path(tmp) / "profile"
        out = Path(tmp) / "out"
        try:
            subprocess.run(
                [
                    soffice,
                    f"-env:UserInstallation=file://{profile}",
                    "--headless",
                    "--calc",
                    "--convert-to",
                    "xlsx:Calc MS Excel 2007 XML",
                    "--outdir",
                    str(out),
                    str(path),
                ],
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
                timeout=timeout,
                check=False,
            )
        except (OSError, subprocess.TimeoutExpired):
            return False
        result = out / Path(path).name
        if not result.exists() or result.stat().st_size == 0:
            return False
        shutil.copyfile(result, path)
    return True


def check_model(original: Path, scrambled: Path, factor: float, tol: float = 1e-4) -> tuple[int, int]:
    """(formula results compared, results that no longer add up). Needs LibreOffice."""
    from openpyxl import load_workbook

    with tempfile.TemporaryDirectory(prefix="dc-check-") as tmp:
        copy = Path(tmp) / "original.xlsx"
        shutil.copyfile(original, copy)
        if not recalc(copy):
            return 0, 0
        real = load_workbook(copy, data_only=True)
        twin_f = load_workbook(scrambled)
        twin_v = load_workbook(scrambled, data_only=True)
        compared = off = 0
        for ws_f, ws_v, ws_r in zip(twin_f.worksheets, twin_v.worksheets, real.worksheets, strict=False):
            for row in ws_f.iter_rows():
                for cell in row:
                    if not (isinstance(cell.value, str) and cell.value.startswith("=")):
                        continue
                    r = ws_r.cell(cell.row, cell.column).value
                    t = ws_v.cell(cell.row, cell.column).value
                    if isinstance(r, bool) or isinstance(t, bool):
                        continue
                    if not isinstance(r, (int, float)) or not isinstance(t, (int, float)):
                        continue
                    compared += 1
                    if not (_close(t, r, tol) or _close(t, r * factor, tol)):
                        off += 1
        return compared, off


def _close(a: float, b: float, tol: float) -> bool:
    return abs(a - b) <= tol * max(1.0, abs(b))
