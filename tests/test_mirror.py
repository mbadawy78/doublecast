"""Mirror folders: scrambled twin, both directions, nothing real leaks into the twin."""

import os

from openpyxl import Workbook, load_workbook

from doublecast.mirror import NOTE, Mirror, register, registered, unregister


def _all_twin_text(twin):
    out = []
    for dirpath, _, files in os.walk(twin):
        for f in files:
            p = os.path.join(dirpath, f)
            out.append(p)
            if p.endswith(".xlsx"):
                wb = load_workbook(p)
                for ws in wb.worksheets:
                    for row in ws.iter_rows(values_only=True):
                        out.extend(str(v) for v in row if v is not None)
            elif not p.endswith((".jpg",)):
                out.append(open(p, encoding="utf-8").read())
    return "\n".join(out)


def _setup(tmp_path):
    real = tmp_path / "work" / "Nordia Hellas deal"
    (real / "Project Falcon").mkdir(parents=True)
    (real / "notes.md").write_text(
        "Call with Eleni Papadopoulou (eleni.p@nordia.gr) about Nordia Hellas pricing.\n", encoding="utf-8"
    )
    (real / "Project Falcon" / "customers.csv").write_text(
        "Customer,Contact,Revenue\nNordia Hellas,Georgios Andreou,1250.00\n", encoding="utf-8"
    )
    wb = Workbook()
    wb.active.append(["Customer", "Margin"])
    wb.active.append(["Kallergis", 0.21])
    wb.save(real / "Kallergis_costing.xlsx")
    (real / "photo.jpg").write_bytes(b"\xff\xd8\xff\xe0 not really a jpeg")
    (real / ".git").mkdir()
    (real / ".git" / "config").write_text("secret", encoding="utf-8")
    return real


def test_twin_holds_no_real_values_and_skips_unreadable_files(engine, tmp_path):
    real = _setup(tmp_path)
    m = Mirror(engine, real)
    assert m.twin.name == "Nordia Hellas deal-AI"
    report = m.sync()
    assert report.scrambled == 3 and report.skipped == {".jpg": 1} and report.failed == 0
    blob = _all_twin_text(m.twin).replace(str(m.twin), "")
    for secret in ("Eleni", "Papadopoulou", "eleni.p@nordia.gr", "Nordia", "Andreou", "Kallergis", "Falcon", "1250"):
        assert secret not in blob, secret
    assert not list(m.twin.rglob("*.jpg")) and not list(m.twin.rglob(".git"))
    assert (m.twin / NOTE).exists()
    # the state with the real names lives in the Doublecast home, not in the twin
    assert m.state_path.parent == engine.config.home / "mirrors"


def test_ai_edit_comes_back_restored(engine, tmp_path):
    real = _setup(tmp_path)
    m = Mirror(engine, real)
    m.sync()
    twin_notes = m.twin / m.state["files"]["notes.md"]["twin"]
    text = twin_notes.read_text(encoding="utf-8")
    double = engine.vault.get("Eleni Papadopoulou")
    assert double and double in text
    twin_notes.write_text(text + f"\nAction: {double} sends the offer.\n", encoding="utf-8")
    report = m.sync()
    assert report.restored == 1
    restored = (real / "notes.md").read_text(encoding="utf-8")
    assert "Action: Eleni Papadopoulou sends the offer." in restored
    assert "eleni.p@nordia.gr" in restored
    # and it does not bounce back and forth
    again = m.sync()
    assert again.scrambled == 0 and again.restored == 0


def test_new_file_from_ai_lands_in_real_folder_with_real_names(engine, tmp_path):
    real = _setup(tmp_path)
    m = Mirror(engine, real)
    m.sync()
    org = engine.vault.get("Nordia Hellas")
    twin_dir = m.twin / m.state["dirs"]["Project Falcon"]
    (twin_dir / f"{org} summary.md").write_text(f"{org} is our biggest customer.", encoding="utf-8")
    m.sync()
    out = real / "Project Falcon" / "Nordia Hellas summary.md"
    assert out.exists()
    assert out.read_text(encoding="utf-8") == "Nordia Hellas is our biggest customer."


def test_both_changed_keeps_both(engine, tmp_path):
    real = _setup(tmp_path)
    m = Mirror(engine, real)
    m.sync()
    twin_notes = m.twin / m.state["files"]["notes.md"]["twin"]
    twin_notes.write_text("AI version", encoding="utf-8")
    (real / "notes.md").write_text("My version, Nordia Hellas.", encoding="utf-8")
    os.utime(real / "notes.md", ns=(1, 1))  # make sure the signature changes
    report = m.sync()
    assert report.conflicts == 1
    assert (real / "notes.md").read_text(encoding="utf-8") == "My version, Nordia Hellas."
    assert (real / "notes (AI).md").read_text(encoding="utf-8") == "AI version"
    assert "Nordia" not in twin_notes.read_text(encoding="utf-8")


def test_real_change_flows_to_twin_and_xlsx_round_trips(engine, tmp_path):
    real = _setup(tmp_path)
    m = Mirror(engine, real)
    m.sync()
    (real / "notes.md").write_text("Georgios Andreou called.", encoding="utf-8")
    assert m.sync().scrambled == 1
    twin_x = m.twin / m.state["files"]["Kallergis_costing.xlsx"]["twin"]
    assert "Kallergis" not in twin_x.name
    wb = load_workbook(twin_x)
    wb.active.append(["New row", 0.3])
    wb.save(twin_x)
    m.sync()
    rows = list(load_workbook(real / "Kallergis_costing.xlsx").active.iter_rows(values_only=True))
    assert rows[1] == ("Kallergis", 0.21)
    # a margin is a ratio: never scaled, so what the AI writes comes back as written
    assert rows[2] == ("New row", 0.3)


def test_twin_inside_real_is_refused(engine, tmp_path):
    real = _setup(tmp_path)
    try:
        Mirror(engine, real, real / "inside")
    except ValueError:
        return
    raise AssertionError("expected ValueError")


def test_registry(tmp_path):
    register(tmp_path, tmp_path / "a", tmp_path / "a-AI")
    register(tmp_path, tmp_path / "a", tmp_path / "a-AI")
    assert len(registered(tmp_path)) == 1
    assert unregister(tmp_path, tmp_path / "a") and registered(tmp_path) == []


def test_interrupted_first_sync_leaves_no_copies_in_real_folder(engine, tmp_path):
    real = _setup(tmp_path)
    m = Mirror(engine, real)
    m.twin.mkdir()
    m._claim()  # a run that started, claimed the twin, and was stopped
    (m.twin / "leftover.md").write_text("half-written by a run that was stopped", encoding="utf-8")
    m.sync()
    m.sync()
    assert not (real / "leftover.md").exists()
    assert not [p for p in real.rglob("*") if "(AI" in p.name]


def test_listed_names_inside_identifiers_and_file_names(engine, tmp_path):
    text = "wb.save('Nordia_Hellas_D365_Items.xlsx')  # Kallergis_costing\n"
    out = engine.scramble(text).text
    assert "Nordia" not in out and "Kallergis" not in out
    assert engine.restore(out) == text


def _model(path, constant_in_formula=False):
    wb = Workbook()
    a = wb.active
    a.title = "Assumptions"
    for row in [
        ["Parameter", "Value"],
        ["Include new lines (1 = yes)", 1],
        ["Cost index", 1.0],
        ["Unit price", 2.5],
        ["Units", 1000],
        ["Days to collect", 60],
        ["Rent per month", 4000],
        ["Material cost per unit", 1.2],
    ]:
        a.append(row)
    p = wb.create_sheet("P&L")
    p.append(["Line", "2026"])
    p.append(["Net sales", "=Assumptions!B4*Assumptions!B5*Assumptions!B2"])
    p.append(["Materials", "=-Assumptions!B8*Assumptions!B5*Assumptions!B3"])
    p.append(["Rent", "=-Assumptions!B7*12" + ("-5000" if constant_in_formula else "")])
    p.append(["EBITDA", "=SUM(B2:B4)"])
    p.append(["EBITDA %", "=B5/B2"])
    p.append(["Receivables", "=B2*Assumptions!B6/365"])
    wb.save(path)


def test_models_keep_adding_up_and_never_show_the_factor(engine, tmp_path):
    import pytest

    from doublecast.files.recalc import find_soffice

    if not find_soffice():
        pytest.skip("LibreOffice not installed")
    real = tmp_path / "Models"
    real.mkdir()
    _model(real / "plan.xlsx")
    _model(real / "broken.xlsx", constant_in_formula=True)
    m = Mirror(engine, real)
    report = m.sync()
    assert report.models_ok == 1 and len(report.models_off) == 1
    twin = load_workbook(m.twin / m.state["files"]["plan.xlsx"]["twin"], data_only=True)
    values = [r[1] for r in twin["Assumptions"].iter_rows(min_row=2, values_only=True)]
    assert values[0] == 1 and values[1] == 1.0 and values[3] == 1000 and values[4] == 60  # never scaled
    assert values[2] != 2.5 and values[5] != 4000  # money is
    pl = [r[1] for r in twin["P&L"].iter_rows(min_row=2, values_only=True)]
    assert all(isinstance(v, (int, float)) for v in pl)  # results filled in
    assert abs(pl[4] - ((2.5 * 1000) - (1.2 * 1000) - 48000) / 2500) < 1e-5  # EBITDA % unchanged (to rounding)
    note = (m.twin / NOTE).read_text(encoding="utf-8")
    assert "Do not rely" in note and m.state["files"]["broken.xlsx"]["twin"] in note


def test_a_second_doublecast_cannot_take_over_a_twin(engine, engine_factory, tmp_path):
    import pytest

    real = _setup(tmp_path)
    Mirror(engine, real).sync()
    other_home = tmp_path / "other"
    other_home.mkdir()
    from doublecast.engine import Engine

    from .conftest import make_config

    other = Engine(make_config(other_home))
    with pytest.raises(RuntimeError, match="another Doublecast"):
        Mirror(other, real).sync()
    other.vault.close()
    assert not [p for p in real.rglob("*") if "(AI" in p.name]


def test_overwritten_files_are_backed_up_and_mass_changes_held(engine, tmp_path):
    real = _setup(tmp_path)
    m = Mirror(engine, real)
    m.sync()
    twin_notes = m.twin / m.state["files"]["notes.md"]["twin"]
    twin_notes.write_text("edited by the AI", encoding="utf-8")
    m.sync()
    backups = list((m.state_path.parent / "backups").rglob("notes.md"))
    assert backups and "Eleni Papadopoulou" in backups[0].read_text(encoding="utf-8")
    for i in range(25):
        (m.twin / f"new{i}.md").write_text("x", encoding="utf-8")
    report = m.sync()
    assert report.held == 25 and not (real / "new0.md").exists()
    assert m.sync(accept=True).restored == 25
