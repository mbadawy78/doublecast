import csv

from docx import Document
from openpyxl import Workbook, load_workbook

from doublecast.files import restore_file, scramble_file


def test_xlsx_round_trip(engine, tmp_path):
    src = tmp_path / "pnl.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Nordia Hellas"
    ws.append(["Customer", "Price", "Volume", "Margin %"])
    ws.append(["Nordia Hellas", 12.5, 100, 0.25])
    ws.append(["Kallergis", 10, 250, 0.31])
    ws["E2"] = "=B2*C2"
    wb.save(src)

    dst, counts, _ = scramble_file(engine, src)
    out = load_workbook(dst)
    s = out.worksheets[0]
    assert "Nordia" not in s.title
    values = [c.value for row in s.iter_rows() for c in row]
    assert "Nordia Hellas" not in values and "Kallergis" not in values
    assert s["B2"].value != 12.5  # price scaled
    assert s["C2"].value == 100  # volume untouched (quantities off by default)
    assert s["D2"].value == 0.25  # percentage untouched
    assert s["E2"].value == "=B2*C2"  # formula kept
    assert counts["MONEY"] == 2

    back, _, _ = restore_file(engine, dst)
    r = load_workbook(back).worksheets[0]
    assert r.title == "Nordia Hellas"
    assert r["A2"].value == "Nordia Hellas"
    assert abs(r["B2"].value - 12.5) < 1e-9
    assert r["B3"].value == 10


def test_docx_keeps_formatting(engine, tmp_path):
    src = tmp_path / "memo.docx"
    doc = Document()
    p = doc.add_paragraph("Meeting with ")
    run = p.add_run("Nordia Hellas")
    run.bold = True
    p.add_run(" about Larissa Plant.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Eleni Papadopoulou"
    table.cell(0, 1).text = "€1.250,00"
    doc.save(src)

    dst, counts, _ = scramble_file(engine, src)
    d2 = Document(str(dst))
    text = d2.paragraphs[0].text
    assert "Nordia" not in text and "Larissa" not in text
    assert d2.paragraphs[0].runs[1].bold  # stand-in kept the bold run
    assert "Papadopoulou" not in d2.tables[0].cell(0, 0).text

    back, _, _ = restore_file(engine, dst)
    d3 = Document(str(back))
    assert d3.paragraphs[0].text == "Meeting with Nordia Hellas about Larissa Plant."
    assert d3.tables[0].cell(0, 1).text == "€1.250,00"


def test_csv_round_trip(engine, tmp_path):
    src = tmp_path / "orders.csv"
    with src.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["customer", "amount", "contact"])
        w.writerow(["Nordia Hellas", "1250.00", "eleni.papadopoulou@acmefoods.gr"])
    dst, _, _ = scramble_file(engine, src)
    text = dst.read_text()
    assert "Nordia" not in text and "1250.00" not in text and "acmefoods" not in text
    back, _, _ = restore_file(engine, dst)
    assert back.read_bytes() == src.read_bytes()


def test_text_file(engine, tmp_path):
    src = tmp_path / "notes.md"
    src.write_bytes("# Notes\r\nCall Georgios Andreou re Project Falcon.\r\n".encode())  # Windows line endings
    dst, _, _ = scramble_file(engine, src)
    assert "Andreou" not in dst.read_text()
    back, _, _ = restore_file(engine, dst)
    assert back.read_bytes() == src.read_bytes()


def test_xlsx_scaled_cells_look_like_money_and_restore_exactly(engine, tmp_path):
    src = tmp_path / "prices.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Item", "Unit price"])
    values = [12.5, 0.4567, 1234567.89, 10, 99.99, 0.01]
    for i, v in enumerate(values):
        ws.append([f"item{i}", v])
    wb.save(src)
    dst, _, _ = scramble_file(engine, src)
    scaled = [r[1] for r in load_workbook(dst).active.iter_rows(min_row=2, values_only=True)]
    assert all(len(repr(float(v)).split(".")[1]) <= 6 for v in scaled)
    assert scaled != values
    back, _, _ = restore_file(engine, dst)
    restored = [r[1] for r in load_workbook(back).active.iter_rows(min_row=2, values_only=True)]
    assert restored == values


def test_spreadsheet_number_restore_is_exact_for_any_factor():
    import random

    from doublecast.files.tabular import scale_value

    rng = random.Random(5)
    values = [12.5, 0.4567, 1234567.89, 10, 99.99, 0.01, 3, 1999.0, 0.1, 250000, 7.25, 0.001, 0.123456]
    for _ in range(5000):
        k = round(rng.uniform(1.17, 3.4), 6)
        for v in values:
            assert scale_value(scale_value(v, k, "scramble"), k, "restore") == v, (v, k)


def test_id_columns_use_the_header_as_context(engine, tmp_path):
    src = tmp_path / "suppliers.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Supplier", "ΑΦΜ", "Τηλέφωνο"])
    ws.append(["Nordia Hellas", 94014201, "2310 123456"])  # ΑΦΜ stored as a number (leading zero lost)
    ws.append(["Kallergis", "094014201", "210 9876543"])
    wb.save(src)
    dst, _, _ = scramble_file(engine, src)
    values = [str(c.value) for row in load_workbook(dst).active.iter_rows(min_row=2) for c in row]
    assert "094014201" not in values and "94014201" not in values
    assert "2310 123456" not in values and "210 9876543" not in values

    csv_src = tmp_path / "suppliers.csv"
    csv_src.write_text("supplier,ΑΦΜ\nNordia Hellas,094014201\n", encoding="utf-8")
    csv_dst, _, _ = scramble_file(engine, csv_src)
    assert "094014201" not in csv_dst.read_text(encoding="utf-8")
    back, _, _ = restore_file(engine, csv_dst)
    assert back.read_text(encoding="utf-8") == csv_src.read_text(encoding="utf-8")


def test_d365_export_aliases_websites_and_contacts_are_scrambled(engine, tmp_path):
    """A real D365 customer export leaked short names, websites and contacts (Oct 2026)."""
    src = tmp_path / "Customers_V3.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Customers_V3"
    ws.append(["CUSTOMERACCOUNT", "NAMEALIAS", "KNOWNAS", "ORGANIZATIONNAME", "PRIMARYCONTACTURL", "Sage price list"])
    ws.append(["C0001", "BRAMBLEX", "Bramblex", "Bramblex Wholesale Canada Ltd.", "www.bramblex.ca", "BRAMBLEX CLUB"])
    ws.append(["C0002", "QUORVIA", "Quorvia", "Quorvia Import Export", "https://quorvia.com/en", "QUORVIA DIST"])
    sheet = wb.create_sheet("Delivery addresses")
    sheet.append(["CUSTOMERACCOUNT", "Customer", "Contact", "City"])
    sheet.append(["C0001", "Bramblex Canada", "Jennifer Walsh", "Ottawa"])
    wb.save(src)
    dst, _, _ = scramble_file(engine, src)
    text = " ".join(str(v) for w in load_workbook(dst).worksheets for r in w.iter_rows(values_only=True) for v in r)
    for real in ("Bramblex", "BRAMBLEX", "bramblex", "Quorvia", "QUORVIA", "quorvia", "Jennifer", "Walsh"):
        assert real not in text, real
    assert "Ottawa" in text  # cities are not secrets
    back, _, _ = restore_file(engine, dst)
    rows = lambda p: [r for w in load_workbook(p).worksheets for r in w.iter_rows(values_only=True)]  # noqa: E731
    assert rows(back) == rows(src)


def test_private_web_addresses_are_swapped_public_ones_kept(engine):
    text = "Order portal: https://www.bramblex.ca/b2b, docs on https://github.com/org/repo"
    out = engine.scramble(text).text
    assert "bramblex" not in out and "https://github.com/org/repo" in out
    assert engine.restore(out) == text


def test_a_real_code_equal_to_an_earlier_double_comes_back_as_itself(engine, tmp_path):
    double = engine.stand_in("ID", "SAGE-11")  # a double made for another file
    src = tmp_path / "items.csv"
    src.write_text(f"Ref,Note\n{double},kept\nSAGE-11,real\n", encoding="utf-8")
    dst, _, _ = scramble_file(engine, src)
    back, _, _ = restore_file(engine, dst)
    assert back.read_text(encoding="utf-8") == src.read_text(encoding="utf-8")


def test_a_real_code_in_a_spreadsheet_equal_to_an_earlier_double(engine, tmp_path):
    double = engine.stand_in("ID", "SAGE-11")
    src = tmp_path / "items.xlsx"
    wb = Workbook()
    wb.active.append(["Ref", "Note"])
    wb.active.append([double, "kept"])
    wb.active.append(["SAGE-11", "real"])
    wb.save(src)
    dst, _, _ = scramble_file(engine, src)
    back, _, _ = restore_file(engine, dst)
    rows = lambda p: list(load_workbook(p).active.iter_rows(values_only=True))  # noqa: E731
    assert rows(back) == rows(src)


def test_codes_given_their_own_double_keep_their_capitals(engine, tmp_path):
    for i in range(30, 60):
        engine.stand_in("ID", f"SAGE-{i}")  # fill the code space with earlier doubles
    src = tmp_path / "codes.csv"
    src.write_text("Ref,Note\n" + "".join(f"SAGE-{i},x\n" for i in range(30, 60)), encoding="utf-8")
    dst, _, _ = scramble_file(engine, src)
    back, _, _ = restore_file(engine, dst)
    assert back.read_text(encoding="utf-8") == src.read_text(encoding="utf-8")
